import streamlit as st
import PyPDF2
import docx
import re
import json
import os
import hashlib
from datetime import datetime

# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Career Risk Analyzer",
    page_icon="📊",
    layout="wide"
)

# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    div[data-testid="stTextInput"] input {
        height: 38px;
        min-height: 38px;
        padding-top: 6px;
        padding-bottom: 6px;
    }

    div[data-testid="stTextArea"] textarea {
        min-height: 90px;
    }

    div[data-testid="column"] div.stButton {
        display: flex;
        justify-content: center;
    }

    div[data-testid="column"] div.stButton > button {
        width: 95%;
        min-height: 45px;
        font-size: 16px;
    }

    div[data-testid="stButton"] > button {
        min-height: 42px;
    }

    div[data-testid="stFormSubmitButton"] > button {
        min-height: 40px;
    }

    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# FILE NAMES
# ============================================================

USER_FILE = "users.json"
PROFILE_FILE = "career_profiles.json"
HISTORY_FILE = "assessment_history.json"

# ============================================================
# SESSION STATE
# ============================================================

if "page" not in st.session_state:
    st.session_state["page"] = "login"

if "username" not in st.session_state:
    st.session_state["username"] = ""

if "navigation_page" not in st.session_state:
    st.session_state["navigation_page"] = "Dashboard"

if "selected_history_index" not in st.session_state:
    st.session_state["selected_history_index"] = None

if "open_roadmap_skill" not in st.session_state:
    st.session_state["open_roadmap_skill"] = None

if "clear_history_confirm" not in st.session_state:
    st.session_state["clear_history_confirm"] = False

if "registration_success" not in st.session_state:
    st.session_state["registration_success"] = False

# ============================================================
# JSON FUNCTIONS
# ============================================================

def load_json(filename):
    if not os.path.exists(filename):
        return {}

    try:
        with open(filename, "r", encoding="utf-8") as file:
            return json.load(file)
    except Exception:
        return {}


def save_json(filename, data):
    with open(filename, "w", encoding="utf-8") as file:
        json.dump(data, file, indent=4)


def load_users():
    return load_json(USER_FILE)


def save_users(users):
    save_json(USER_FILE, users)


def load_profiles():
    return load_json(PROFILE_FILE)


def save_profiles(profiles):
    save_json(PROFILE_FILE, profiles)


def load_history():
    return load_json(HISTORY_FILE)


def save_history(history):
    save_json(HISTORY_FILE, history)

# ============================================================
# CORE LANGUAGES
# ============================================================

CORE_LANGUAGES = [
    "Python",
    "Java",
    "C++"
]

# ============================================================
# ROLE SKILLS
# ============================================================

ROLE_SKILLS = {
    "Data Analyst": [
        "Python",
        "SQL",
        "Excel",
        "Power BI",
        "Tableau",
        "Statistics",
        "Pandas",
        "Data Visualization"
    ],

    "Software Developer": [
        "Python",
        "Java",
        "Data Structures",
        "Algorithms",
        "OOP",
        "SQL",
        "Git"
    ],

    "Data Scientist": [
        "Python",
        "SQL",
        "Statistics",
        "Machine Learning",
        "Pandas",
        "NumPy",
        "Data Visualization"
    ],

    "Machine Learning Engineer": [
        "Python",
        "Machine Learning",
        "Scikit-learn",
        "TensorFlow",
        "Deep Learning",
        "SQL",
        "Deployment"
    ],

    "Web Developer": [
        "HTML",
        "CSS",
        "JavaScript",
        "React",
        "Flask",
        "Django",
        "SQL",
        "Git"
    ]
}

# ============================================================
# RECOMMENDATIONS
# ============================================================

RECOMMENDATIONS = {
    "Python":
        "Practice Python programming, functions, OOP, data structures and problem solving.",

    "SQL":
        "Learn SELECT, JOIN, GROUP BY, subqueries and aggregate functions.",

    "Excel":
        "Learn formulas, lookup functions, pivot tables and data cleaning.",

    "Power BI":
        "Learn Power Query, data modelling, DAX and dashboard creation.",

    "Tableau":
        "Learn Tableau charts, dashboards, filters and calculated fields.",

    "Statistics":
        "Study descriptive statistics, probability, distributions and hypothesis testing.",

    "Pandas":
        "Practice data cleaning, filtering, grouping and analysis using Pandas.",

    "Data Visualization":
        "Learn how to create meaningful charts and communicate insights.",

    "Java":
        "Practice Java syntax, OOP, collections and exception handling.",

    "Data Structures":
        "Learn arrays, lists, stacks, queues, trees, graphs and hash tables.",

    "Algorithms":
        "Practice searching, sorting, recursion and algorithmic problem solving.",

    "OOP":
        "Learn classes, objects, inheritance, polymorphism and encapsulation.",

    "Git":
        "Learn repositories, commits, branches, merge and GitHub workflows.",

    "Machine Learning":
        "Study supervised and unsupervised learning and model evaluation.",

    "NumPy":
        "Practice arrays, indexing, mathematical operations and numerical computing.",

    "Scikit-learn":
        "Learn preprocessing, model training, evaluation and pipelines.",

    "TensorFlow":
        "Learn tensors, neural networks, model training and evaluation.",

    "Deep Learning":
        "Study neural networks, CNNs, RNNs and deep learning concepts.",

    "Deployment":
        "Learn how to deploy Python and machine learning applications.",

    "HTML":
        "Learn HTML elements, forms, semantic tags and page structure.",

    "CSS":
        "Learn selectors, layouts, Flexbox, Grid and responsive design.",

    "JavaScript":
        "Learn variables, functions, DOM manipulation and asynchronous programming.",

    "React":
        "Learn components, props, state, hooks and API integration.",

    "Flask":
        "Learn Flask routes, templates, forms and APIs.",

    "Django":
        "Learn Django projects, apps, models, views, URLs and templates."
}

# ============================================================
# LEARNING RESOURCES
# ============================================================

LEARNING_RESOURCES = {
    "Python": "https://docs.python.org/3/tutorial/",
    "SQL": "https://www.postgresql.org/docs/current/tutorial.html",
    "Excel": "https://support.microsoft.com/excel",
    "Power BI": "https://learn.microsoft.com/en-us/training/paths/prepare-data-power-bi/",
    "Tableau": "https://www.tableau.com/learn/training",
    "Statistics": "https://www.khanacademy.org/math/statistics-probability",
    "Pandas": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/index.html",
    "Data Visualization": "https://matplotlib.org/stable/users/explain/quick_start.html",
    "Java": "https://dev.java/learn/",
    "Data Structures": "https://www.geeksforgeeks.org/dsa/",
    "Algorithms": "https://www.geeksforgeeks.org/fundamentals-of-algorithms/",
    "OOP": "https://docs.python.org/3/tutorial/classes.html",
    "Git": "https://git-scm.com/book/en/v2",
    "Machine Learning": "https://developers.google.com/machine-learning/intro",
    "NumPy": "https://numpy.org/doc/stable/user/absolute_beginners.html",
    "Scikit-learn": "https://scikit-learn.org/stable/user_guide.html",
    "TensorFlow": "https://www.tensorflow.org/learn",
    "Deep Learning": "https://www.deeplearning.ai/courses/",
    "Deployment": "https://docs.streamlit.io/deploy",
    "HTML": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content",
    "CSS": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Styling_basics",
    "JavaScript": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Scripting",
    "React": "https://react.dev/learn",
    "Flask": "https://flask.palletsprojects.com/en/stable/quickstart/",
    "Django": "https://docs.djangoproject.com/en/stable/intro/tutorial01/"
}

# ============================================================
# SKILL DEPENDENCIES
# ============================================================

SKILL_DEPENDENCIES = {
    "Excel": [],
    "SQL": [],
    "Power BI": ["Excel", "SQL"],
    "Tableau": ["SQL"],
    "Statistics": [],
    "Pandas": ["Python"],
    "Data Visualization": ["Pandas"],
    "Data Structures": ["Python"],
    "Algorithms": ["Data Structures"],
    "OOP": ["Python"],
    "Git": [],
    "NumPy": ["Python"],

    "Machine Learning": [
        "Python",
        "Statistics",
        "Pandas"
    ],

    "Scikit-learn": [
        "Python",
        "Machine Learning"
    ],

    "TensorFlow": [
        "Python",
        "Machine Learning"
    ],

    "Deep Learning": [
        "Python",
        "Machine Learning"
    ],

    "Deployment": [
        "Python"
    ],

    "HTML": [],
    "CSS": ["HTML"],

    "JavaScript": [
        "HTML",
        "CSS"
    ],

    "React": [
        "JavaScript"
    ],

    "Flask": [
        "Python"
    ],

    "Django": [
        "Python"
    ]
}

# ============================================================
# DETAILED ROADMAP
# ============================================================

ROADMAP_DETAILS = {

    # ========================================================
    # PYTHON
    # ========================================================

    "Python": {
        "concepts": [
            {
                "name": "Python Basics, Variables and Data Types",
                "link": "https://docs.python.org/3/tutorial/introduction.html"
            },
            {
                "name": "Control Statements",
                "link": "https://www.geeksforgeeks.org/computer-science-fundamentals/control-flow-statements-in-programming/"
            },
            {
                "name": "Functions",
                "link": "https://www.geeksforgeeks.org/python/python-functions/"
            },
            {
                "name": "Lists and List Comprehensions",
                "link": "https://docs.python.org/3/tutorial/datastructures.html"
            },
            {
                "name": "Tuples, Sets and Dictionaries",
                "link": "https://docs.python.org/3/tutorial/datastructures.html"
            },
            {
                "name": "Modules and Packages",
                "link": "https://docs.python.org/3/tutorial/modules.html"
            },
            {
                "name": "File Handling",
                "link": "https://docs.python.org/3/tutorial/inputoutput.html#reading-and-writing-files"
            },
            {
                "name": "Exception Handling",
                "link": "https://docs.python.org/3/tutorial/errors.html"
            },
            {
                "name": "Object-Oriented Programming",
                "link": "https://docs.python.org/3/tutorial/classes.html"
            }
        ],
        "practice":
            "Solve programming problems involving conditions, loops, functions, collections, files and OOP.",
        "project":
            "Build a real-world Employee Management System with employee records, search, update, delete and persistent data storage."
    },

    # ========================================================
    # SQL
    # ========================================================

    "SQL": {
        "concepts": [
            {
                "name": "SELECT and FROM",
                "link": "https://www.postgresql.org/docs/current/tutorial-select.html"
            },
            {
                "name": "WHERE and Filtering",
                "link": "https://www.postgresql.org/docs/current/tutorial-select.html"
            },
            {
                "name": "ORDER BY and DISTINCT",
                "link": "https://www.postgresql.org/docs/current/tutorial-select.html"
            },
            {
                "name": "Aggregate Functions",
                "link": "https://www.postgresql.org/docs/current/tutorial-agg.html"
            },
            {
                "name": "GROUP BY and HAVING",
                "link": "https://www.postgresql.org/docs/current/tutorial-agg.html"
            },
            {
                "name": "INNER JOIN and OUTER JOIN",
                "link": "https://www.postgresql.org/docs/current/tutorial-join.html"
            },
            {
                "name": "Subqueries",
                "link": "https://www.postgresql.org/docs/current/functions-subquery.html"
            },
            {
                "name": "CASE Expressions",
                "link": "https://www.postgresql.org/docs/current/functions-conditional.html"
            },
            {
                "name": "Common Table Expressions",
                "link": "https://www.postgresql.org/docs/current/queries-with.html"
            },
            {
                "name": "Window Functions",
                "link": "https://www.postgresql.org/docs/current/tutorial-window.html"
            }
        ],
        "practice":
            "Write queries using realistic employee, customer, product and sales databases.",
        "project":
            "Build a Sales Analytics Database and create SQL reports for revenue, customers, products and regional performance."
    },

    # ========================================================
    # EXCEL
    # ========================================================

    "Excel": {
        "concepts": [
            {
                "name": "Excel Fundamentals",
                "link": "https://support.microsoft.com/excel"
            },
            {
                "name": "Formulas and Functions",
                "link": "https://support.microsoft.com/office/overview-of-formulas-in-excel-ecfdc708-9162-49e8-b993-c311f47ca173"
            },
            {
                "name": "IF Function",
                "link": "https://support.microsoft.com/office/if-function-69aed7c9-4e8a-4755-a9bc-aa8bbff73be2"
            },
            {
                "name": "XLOOKUP",
                "link": "https://support.microsoft.com/office/xlookup-function-b7fd680e-6d10-43e6-84f9-88eae8bf5929"
            },
            {
                "name": "Sorting and Filtering",
                "link": "https://support.microsoft.com/excel"
            },
            {
                "name": "Pivot Tables",
                "link": "https://support.microsoft.com/office/create-a-pivottable-to-analyze-worksheet-data-a9a84538-bfe9-40a9-a8e9-f99134456576"
            },
            {
                "name": "Charts",
                "link": "https://support.microsoft.com/excel"
            },
            {
                "name": "Conditional Formatting",
                "link": "https://support.microsoft.com/excel"
            }
        ],
        "practice":
            "Clean a real sales dataset and create formulas, lookup operations, pivot tables and charts.",
        "project":
            "Build a real-world Sales Performance Dashboard showing revenue, profit, products, regions and monthly trends."
    },

    # ========================================================
    # POWER BI
    # ========================================================

    "Power BI": {
        "concepts": [
            {
                "name": "Power BI Desktop Basics",
                "link": "https://learn.microsoft.com/en-us/training/modules/get-started-with-power-bi/"
            },
            {
                "name": "Getting Data into Power BI",
                "link": "https://learn.microsoft.com/en-us/training/modules/get-data/"
            },
            {
                "name": "Power Query",
                "link": "https://learn.microsoft.com/en-us/power-query/power-query-what-is-power-query"
            },
            {
                "name": "Data Cleaning and Transformation",
                "link": "https://learn.microsoft.com/en-us/training/paths/prepare-data-power-bi/"
            },
            {
                "name": "Data Relationships",
                "link": "https://learn.microsoft.com/en-us/power-bi/transform-model/desktop-relationships-understand"
            },
            {
                "name": "Star Schema and Data Modelling",
                "link": "https://learn.microsoft.com/en-us/power-bi/guidance/star-schema"
            },
            {
                "name": "DAX Fundamentals",
                "link": "https://learn.microsoft.com/en-us/dax/"
            },
            {
                "name": "Measures",
                "link": "https://learn.microsoft.com/en-us/power-bi/transform-model/desktop-measures"
            },
            {
                "name": "Reports and Dashboards",
                "link": "https://learn.microsoft.com/en-us/power-bi/create-reports/"
            }
        ],
        "practice":
            "Import a real dataset, clean it using Power Query, create relationships and build an interactive report.",
        "project":
            "Build a real-world Business Sales Dashboard showing revenue, profit, customers, products, regions and monthly trends."
    },

    # ========================================================
    # TABLEAU
    # ========================================================

    "Tableau": {
        "concepts": [
            {
                "name": "Tableau Getting Started",
                "link": "https://www.tableau.com/learn"
            },
            {
                "name": "Connecting to Data",
                "link": "https://www.tableau.com/learn/training"
            },
            {
                "name": "Dimensions and Measures",
                "link": "https://www.tableau.com/learn/training"
            },
            {
                "name": "Charts and Visualizations",
                "link": "https://www.tableau.com/learn/training"
            },
            {
                "name": "Filters",
                "link": "https://help.tableau.com/current/pro/desktop/en-us/filtering.htm"
            },
            {
                "name": "Calculated Fields",
                "link": "https://help.tableau.com/current/pro/desktop/en-us/calculations_calculatedfields_create.htm"
            },
            {
                "name": "Dashboards",
                "link": "https://help.tableau.com/current/pro/desktop/en-us/dashboards.htm"
            }
        ],
        "practice":
            "Connect Tableau to a real dataset and create interactive business visualizations.",
        "project":
            "Build a real-world Tableau Business Intelligence Dashboard for sales and customer performance."
    },

    # ========================================================
    # STATISTICS
    # ========================================================

    "Statistics": {
        "concepts": [
            {
                "name": "Mean, Median and Mode",
                "link": "https://www.khanacademy.org/math/statistics-probability/summarizing-quantitative-data"
            },
            {
                "name": "Variance and Standard Deviation",
                "link": "https://www.khanacademy.org/math/statistics-probability"
            },
            {
                "name": "Probability",
                "link": "https://www.khanacademy.org/math/statistics-probability/probability-library"
            },
            {
                "name": "Probability Distributions",
                "link": "https://www.khanacademy.org/math/statistics-probability/random-variables-stats-library"
            },
            {
                "name": "Correlation",
                "link": "https://www.khanacademy.org/math/statistics-probability/describing-relationships-quantitative-data"
            },
            {
                "name": "Hypothesis Testing",
                "link": "https://www.khanacademy.org/math/statistics-probability/significance-tests-one-sample"
            },
            {
                "name": "Regression",
                "link": "https://www.khanacademy.org/math/statistics-probability"
            }
        ],
        "practice":
            "Analyse real datasets and calculate descriptive statistics, correlations and hypothesis tests.",
        "project":
            "Perform statistical analysis of customer behaviour or sales performance using a real-world dataset."
    },

    # ========================================================
    # PANDAS
    # ========================================================

    "Pandas": {
        "concepts": [
            {
                "name": "Pandas Introduction",
                "link": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/index.html"
            },
            {
                "name": "Series and DataFrames",
                "link": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/01_table_oriented.html"
            },
            {
                "name": "Reading CSV and Excel Files",
                "link": "https://pandas.pydata.org/docs/user_guide/io.html"
            },
            {
                "name": "Selecting and Filtering Data",
                "link": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/03_subset_data.html"
            },
            {
                "name": "Missing Values",
                "link": "https://pandas.pydata.org/docs/user_guide/missing_data.html"
            },
            {
                "name": "GroupBy",
                "link": "https://pandas.pydata.org/docs/user_guide/groupby.html"
            },
            {
                "name": "Merge and Join",
                "link": "https://pandas.pydata.org/docs/user_guide/merging.html"
            },
            {
                "name": "Sorting and Aggregation",
                "link": "https://pandas.pydata.org/docs/getting_started/intro_tutorials/index.html"
            }
        ],
        "practice":
            "Clean, transform and analyse real CSV datasets using Pandas.",
        "project":
            "Build an end-to-end Customer Sales Data Analysis project using Pandas."
    },

    # ========================================================
    # DATA VISUALIZATION
    # ========================================================

    "Data Visualization": {
        "concepts": [
            {
                "name": "Matplotlib Basics",
                "link": "https://matplotlib.org/stable/users/explain/quick_start.html"
            },
            {
                "name": "Line Charts",
                "link": "https://matplotlib.org/stable/api/_as_gen/matplotlib.pyplot.plot.html"
            },
            {
                "name": "Bar Charts",
                "link": "https://matplotlib.org/stable/api/_as_gen/matplotlib.pyplot.bar.html"
            },
            {
                "name": "Scatter Plots",
                "link": "https://matplotlib.org/stable/api/_as_gen/matplotlib.pyplot.scatter.html"
            },
            {
                "name": "Histograms",
                "link": "https://matplotlib.org/stable/api/_as_gen/matplotlib.pyplot.hist.html"
            },
            {
                "name": "Data Visualization Principles",
                "link": "https://www.tableau.com/learn/articles/data-visualization"
            }
        ],
        "practice":
            "Create charts from real datasets and explain the business insight shown by each visualization.",
        "project":
            "Build a Business Analytics Dashboard showing sales, customer and performance insights."
    },

    # ========================================================
    # JAVA
    # ========================================================

    "Java": {
        "concepts": [
            {
                "name": "Java Language Basics",
                "link": "https://dev.java/learn/language-basics/"
            },
            {
                "name": "Variables and Data Types",
                "link": "https://dev.java/learn/language-basics/"
            },
            {
                "name": "Conditions and Loops",
                "link": "https://dev.java/learn/language-basics/"
            },
            {
                "name": "Classes and Objects",
                "link": "https://dev.java/learn/classes-objects/"
            },
            {
                "name": "Inheritance",
                "link": "https://dev.java/learn/inheritance/"
            },
            {
                "name": "Interfaces",
                "link": "https://dev.java/learn/interfaces/"
            },
            {
                "name": "Collections",
                "link": "https://dev.java/learn/api/collections-framework/"
            },
            {
                "name": "Exception Handling",
                "link": "https://dev.java/learn/exceptions/"
            }
        ],
        "practice":
            "Solve Java programming problems and build small console applications.",
        "project":
            "Build a Java Employee Management System with CRUD operations and persistent storage."
    },

    # ========================================================
    # DATA STRUCTURES
    # ========================================================

    "Data Structures": {
        "concepts": [
            {
                "name": "Arrays",
                "link": "https://www.geeksforgeeks.org/dsa/array-data-structure/"
            },
            {
                "name": "Linked Lists",
                "link": "https://www.geeksforgeeks.org/dsa/linked-list-data-structure/"
            },
            {
                "name": "Stacks",
                "link": "https://www.geeksforgeeks.org/dsa/stack-data-structure/"
            },
            {
                "name": "Queues",
                "link": "https://www.geeksforgeeks.org/dsa/queue-data-structure/"
            },
            {
                "name": "Hash Tables",
                "link": "https://www.geeksforgeeks.org/dsa/hashing-data-structure/"
            },
            {
                "name": "Trees",
                "link": "https://www.geeksforgeeks.org/dsa/binary-tree-data-structure/"
            },
            {
                "name": "Graphs",
                "link": "https://www.geeksforgeeks.org/dsa/graph-data-structure-and-algorithms/"
            }
        ],
        "practice":
            "Solve coding problems using arrays, linked lists, stacks, queues, trees and graphs.",
        "project":
            "Build an Inventory Management System using suitable data structures for searching and storing products."
    },

    # ========================================================
    # ALGORITHMS
    # ========================================================

    "Algorithms": {
        "concepts": [
            {
                "name": "Time and Space Complexity",
                "link": "https://www.geeksforgeeks.org/dsa/understanding-time-complexity-simple-examples/"
            },
            {
                "name": "Linear Search",
                "link": "https://www.geeksforgeeks.org/dsa/linear-search/"
            },
            {
                "name": "Binary Search",
                "link": "https://www.geeksforgeeks.org/dsa/binary-search/"
            },
            {
                "name": "Sorting Algorithms",
                "link": "https://www.geeksforgeeks.org/dsa/sorting-algorithms/"
            },
            {
                "name": "Recursion",
                "link": "https://www.geeksforgeeks.org/dsa/recursion/"
            },
            {
                "name": "Greedy Algorithms",
                "link": "https://www.geeksforgeeks.org/dsa/greedy-algorithms/"
            },
            {
                "name": "Dynamic Programming",
                "link": "https://www.geeksforgeeks.org/dsa/dynamic-programming/"
            }
        ],
        "practice":
            "Solve progressively harder algorithmic problems and analyse time and space complexity.",
        "project":
            "Build an application that compares searching and sorting algorithms using real datasets."
    },

    # ========================================================
    # OOP
    # ========================================================

    "OOP": {
        "concepts": [
            {
                "name": "Classes and Objects",
                "link": "https://docs.python.org/3/tutorial/classes.html"
            },
            {
                "name": "Constructors",
                "link": "https://docs.python.org/3/tutorial/classes.html"
            },
            {
                "name": "Encapsulation",
                "link": "https://docs.python.org/3/tutorial/classes.html"
            },
            {
                "name": "Inheritance",
                "link": "https://docs.python.org/3/tutorial/classes.html#inheritance"
            },
            {
                "name": "Polymorphism",
                "link": "https://docs.python.org/3/tutorial/classes.html#inheritance"
            },
            {
                "name": "Abstraction",
                "link": "https://docs.python.org/3/library/abc.html"
            }
        ],
        "practice":
            "Create small applications using classes, inheritance, encapsulation and polymorphism.",
        "project":
            "Build a real-world Library Management System using object-oriented design."
    },

    # ========================================================
    # GIT
    # ========================================================

    "Git": {
        "concepts": [
            {
                "name": "Git Installation",
                "link": "https://git-scm.com/book/en/v2/Getting-Started-Installing-Git"
            },
            {
                "name": "Creating a Repository",
                "link": "https://git-scm.com/book/en/v2/Git-Basics-Getting-a-Git-Repository"
            },
            {
                "name": "Git Add and Commit",
                "link": "https://git-scm.com/docs/git-commit"
            },
            {
                "name": "Branches",
                "link": "https://git-scm.com/book/en/v2/Git-Branching-Branches-in-a-Nutshell"
            },
            {
                "name": "Merging",
                "link": "https://git-scm.com/book/en/v2/Git-Branching-Basic-Branching-and-Merging"
            },
            {
                "name": "Remote Repositories",
                "link": "https://git-scm.com/book/en/v2/Git-Basics-Working-with-Remotes"
            },
            {
                "name": "GitHub Pull Requests",
                "link": "https://docs.github.com/en/pull-requests"
            }
        ],
        "practice":
            "Use Git for your projects with meaningful commits, branches and remote repositories.",
        "project":
            "Create and manage a complete software project on GitHub using branches and pull requests."
    },

    # ========================================================
    # NUMPY
    # ========================================================

    "NumPy": {
        "concepts": [
            {
                "name": "NumPy Introduction",
                "link": "https://numpy.org/doc/stable/user/absolute_beginners.html"
            },
            {
                "name": "NumPy Arrays",
                "link": "https://numpy.org/doc/stable/user/basics.creation.html"
            },
            {
                "name": "Indexing and Slicing",
                "link": "https://numpy.org/doc/stable/user/basics.indexing.html"
            },
            {
                "name": "Array Shape and Reshaping",
                "link": "https://numpy.org/doc/stable/user/quickstart.html"
            },
            {
                "name": "Mathematical Operations",
                "link": "https://numpy.org/doc/stable/user/quickstart.html"
            },
            {
                "name": "Broadcasting",
                "link": "https://numpy.org/doc/stable/user/basics.broadcasting.html"
            },
            {
                "name": "Linear Algebra",
                "link": "https://numpy.org/doc/stable/reference/routines.linalg.html"
            }
        ],
        "practice":
            "Perform numerical calculations and manipulate real datasets using NumPy.",
        "project":
            "Build a numerical data-analysis project using NumPy for statistical and mathematical calculations."
    },

    # ========================================================
    # MACHINE LEARNING
    # ========================================================

    "Machine Learning": {
        "concepts": [
            {
                "name": "Machine Learning Fundamentals",
                "link": "https://developers.google.com/machine-learning/intro"
            },
            {
                "name": "Supervised Learning",
                "link": "https://developers.google.com/machine-learning/intro-to-ml/supervised-learning"
            },
            {
                "name": "Unsupervised Learning",
                "link": "https://developers.google.com/machine-learning/clustering/overview"
            },
            {
                "name": "Training and Testing Data",
                "link": "https://scikit-learn.org/stable/modules/cross_validation.html"
            },
            {
                "name": "Feature Engineering",
                "link": "https://scikit-learn.org/stable/modules/preprocessing.html"
            },
            {
                "name": "Classification",
                "link": "https://scikit-learn.org/stable/supervised_learning.html"
            },
            {
                "name": "Regression",
                "link": "https://scikit-learn.org/stable/supervised_learning.html"
            },
            {
                "name": "Model Evaluation",
                "link": "https://scikit-learn.org/stable/modules/model_evaluation.html"
            },
            {
                "name": "Overfitting and Underfitting",
                "link": "https://developers.google.com/machine-learning/crash-course/overfitting/overfitting"
            }
        ],
        "practice":
            "Train classification and regression models on real datasets and evaluate their performance.",
        "project":
            "Build an end-to-end machine-learning prediction system using a public dataset, preprocessing, model training and evaluation."
    },

    # ========================================================
    # SCIKIT-LEARN
    # ========================================================

    "Scikit-learn": {
        "concepts": [
            {
                "name": "Preprocessing",
                "link": "https://scikit-learn.org/stable/modules/preprocessing.html"
            },
            {
                "name": "Train-Test Split",
                "link": "https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.train_test_split.html"
            },
            {
                "name": "Feature Scaling",
                "link": "https://scikit-learn.org/stable/modules/preprocessing.html"
            },
            {
                "name": "Classification",
                "link": "https://scikit-learn.org/stable/supervised_learning.html"
            },
            {
                "name": "Regression",
                "link": "https://scikit-learn.org/stable/supervised_learning.html"
            },
            {
                "name": "Cross Validation",
                "link": "https://scikit-learn.org/stable/modules/cross_validation.html"
            },
            {
                "name": "Model Evaluation",
                "link": "https://scikit-learn.org/stable/modules/model_evaluation.html"
            },
            {
                "name": "Pipelines",
                "link": "https://scikit-learn.org/stable/modules/compose.html"
            }
        ],
        "practice":
            "Train several machine-learning models, compare metrics and build reusable preprocessing pipelines.",
        "project":
            "Build an end-to-end Scikit-learn prediction application with preprocessing, model training, evaluation and prediction."
    },

    # ========================================================
    # TENSORFLOW
    # ========================================================

    "TensorFlow": {
        "concepts": [
            {
                "name": "Tensor Basics",
                "link": "https://www.tensorflow.org/guide/tensor"
            },
            {
                "name": "TensorFlow Data",
                "link": "https://www.tensorflow.org/guide/data"
            },
            {
                "name": "Keras Sequential Models",
                "link": "https://www.tensorflow.org/guide/keras/sequential_model"
            },
            {
                "name": "Neural Network Layers",
                "link": "https://www.tensorflow.org/guide/keras/sequential_model"
            },
            {
                "name": "Training a Model",
                "link": "https://www.tensorflow.org/guide/keras/training_with_built_in_methods"
            },
            {
                "name": "Model Evaluation",
                "link": "https://www.tensorflow.org/guide/keras/training_with_built_in_methods"
            }
        ],
        "practice":
            "Build and train neural networks using TensorFlow and Keras.",
        "project":
            "Build a real-world image or text classification application using TensorFlow."
    },

    # ========================================================
    # DEEP LEARNING
    # ========================================================

    "Deep Learning": {
        "concepts": [
            {
                "name": "Neural Network Fundamentals",
                "link": "https://www.deeplearning.ai/courses/neural-networks-deep-learning/"
            },
            {
                "name": "Activation Functions",
                "link": "https://developers.google.com/machine-learning/crash-course/neural-networks/activation-functions"
            },
            {
                "name": "Gradient Descent",
                "link": "https://developers.google.com/machine-learning/crash-course/reducing-loss/gradient-descent"
            },
            {
                "name": "Backpropagation",
                "link": "https://developers.google.com/machine-learning/crash-course/neural-networks/backpropagation"
            },
            {
                "name": "Convolutional Neural Networks",
                "link": "https://www.tensorflow.org/tutorials/images/cnn"
            },
            {
                "name": "Recurrent Neural Networks",
                "link": "https://www.tensorflow.org/guide/keras/working_with_rnns"
            },
            {
                "name": "Model Evaluation",
                "link": "https://www.tensorflow.org/guide/keras/training_with_built_in_methods"
            }
        ],
        "practice":
            "Train neural networks using image or text datasets.",
        "project":
            "Build a real-world image classification or text classification system using deep learning."
    },

    # ========================================================
    # DEPLOYMENT
    # ========================================================

    "Deployment": {
        "concepts": [
            {
                "name": "Python Project Structure",
                "link": "https://docs.python.org/3/tutorial/modules.html"
            },
            {
                "name": "Installing Dependencies",
                "link": "https://packaging.python.org/en/latest/tutorials/installing-packages/"
            },
            {
                "name": "Requirements File",
                "link": "https://pip.pypa.io/en/stable/reference/requirements-file-format/"
            },
            {
                "name": "Environment Variables",
                "link": "https://12factor.net/config"
            },
            {
                "name": "REST API Fundamentals",
                "link": "https://developer.mozilla.org/en-US/docs/Glossary/REST"
            },
            {
                "name": "Streamlit Deployment",
                "link": "https://docs.streamlit.io/deploy"
            }
        ],
        "practice":
            "Package a Python application, create its requirements file and deploy it so other users can access it.",
        "project":
            "Deploy a real-world machine-learning or Streamlit application and make it accessible through the internet."
    },

    # ========================================================
    # HTML
    # ========================================================

    "HTML": {
        "concepts": [
            {
                "name": "HTML Document Structure",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content/Basic_HTML_syntax"
            },
            {
                "name": "Headings and Paragraphs",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content"
            },
            {
                "name": "Links",
                "link": "https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Elements/a"
            },
            {
                "name": "Images",
                "link": "https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Elements/img"
            },
            {
                "name": "Lists and Tables",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content"
            },
            {
                "name": "Forms and Input",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Extensions/Forms"
            },
            {
                "name": "Semantic HTML",
                "link": "https://developer.mozilla.org/en-US/docs/Glossary/Semantics"
            }
        ],
        "practice":
            "Create web pages using semantic HTML and properly structured forms.",
        "project":
            "Build a real-world personal portfolio and job application website."
    },

    # ========================================================
    # CSS
    # ========================================================

    "CSS": {
        "concepts": [
            {
                "name": "CSS Basics and Selectors",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Styling_basics"
            },
            {
                "name": "Box Model",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Styling_basics/Box_model"
            },
            {
                "name": "Typography",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Text_styling"
            },
            {
                "name": "Flexbox",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/CSS_layout/Flexbox"
            },
            {
                "name": "CSS Grid",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/CSS_layout/Grids"
            },
            {
                "name": "Responsive Web Design",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/CSS_layout/Responsive_Design"
            }
        ],
        "practice":
            "Convert static HTML pages into responsive interfaces that work on desktop and mobile.",
        "project":
            "Build a responsive portfolio or company landing page."
    },

    # ========================================================
    # JAVASCRIPT
    # ========================================================

    "JavaScript": {
        "concepts": [
            {
                "name": "Variables and Data Types",
                "link": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide/Grammar_and_types"
            },
            {
                "name": "Functions",
                "link": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide/Functions"
            },
            {
                "name": "Arrays",
                "link": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Reference/Global_Objects/Array"
            },
            {
                "name": "Objects",
                "link": "https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide/Working_with_objects"
            },
            {
                "name": "DOM",
                "link": "https://developer.mozilla.org/en-US/docs/Web/API/Document_Object_Model/Introduction"
            },
            {
                "name": "Events",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Scripting/Events"
            },
            {
                "name": "Asynchronous JavaScript",
                "link": "https://developer.mozilla.org/en-US/docs/Learn_web_development/Extensions/Async_JS"
            },
            {
                "name": "Fetch API",
                "link": "https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API"
            }
        ],
        "practice":
            "Create interactive web pages and consume APIs using JavaScript.",
        "project":
            "Build a real-world task management application with frontend interaction and API integration."
    },

    # ========================================================
    # REACT
    # ========================================================

    "React": {
        "concepts": [
            {
                "name": "React Components",
                "link": "https://react.dev/learn/your-first-component"
            },
            {
                "name": "JSX",
                "link": "https://react.dev/learn/writing-markup-with-jsx"
            },
            {
                "name": "Props",
                "link": "https://react.dev/learn/passing-props-to-a-component"
            },
            {
                "name": "State",
                "link": "https://react.dev/learn/state-a-components-memory"
            },
            {
                "name": "Event Handling",
                "link": "https://react.dev/learn/responding-to-events"
            },
            {
                "name": "Hooks",
                "link": "https://react.dev/reference/react"
            },
            {
                "name": "Using APIs",
                "link": "https://react.dev/learn/synchronizing-with-effects"
            }
        ],
        "practice":
            "Build reusable React components and connect the frontend to APIs.",
        "project":
            "Build a real-world Job Portal frontend with login, job search, job details and application functionality."
    },

    # ========================================================
    # FLASK
    # ========================================================

    "Flask": {
        "concepts": [
            {
                "name": "Flask Application Basics",
                "link": "https://flask.palletsprojects.com/en/stable/quickstart/"
            },
            {
                "name": "Routing",
                "link": "https://flask.palletsprojects.com/en/stable/quickstart/#routing"
            },
            {
                "name": "Templates",
                "link": "https://flask.palletsprojects.com/en/stable/templating/"
            },
            {
                "name": "Request and Form Data",
                "link": "https://flask.palletsprojects.com/en/stable/quickstart/#the-request-object"
            },
            {
                "name": "JSON APIs",
                "link": "https://flask.palletsprojects.com/en/stable/api/"
            },
            {
                "name": "Database Integration",
                "link": "https://flask.palletsprojects.com/en/stable/tutorial/database/"
            }
        ],
        "practice":
            "Create Flask routes, forms and APIs and connect the application to a database.",
        "project":
            "Build a real-world Employee Management System using Flask and a relational database."
    },

    # ========================================================
    # DJANGO
    # ========================================================

    "Django": {
        "concepts": [
            {
                "name": "Django Project and App",
                "link": "https://docs.djangoproject.com/en/stable/intro/tutorial01/"
            },
            {
                "name": "URLs",
                "link": "https://docs.djangoproject.com/en/stable/topics/http/urls/"
            },
            {
                "name": "Views",
                "link": "https://docs.djangoproject.com/en/stable/topics/http/views/"
            },
            {
                "name": "Templates",
                "link": "https://docs.djangoproject.com/en/stable/topics/templates/"
            },
            {
                "name": "Models",
                "link": "https://docs.djangoproject.com/en/stable/topics/db/models/"
            },
            {
                "name": "Forms",
                "link": "https://docs.djangoproject.com/en/stable/topics/forms/"
            },
            {
                "name": "Authentication",
                "link": "https://docs.djangoproject.com/en/stable/topics/auth/"
            }
        ],
        "practice":
            "Build Django applications using models, views, templates, forms and authentication.",
        "project":
            "Build a real-world Job Portal or Employee Management System using Django and a relational database."
    }
}

# ============================================================
# PASSWORD VALIDATION
# ============================================================

def validate_password(password):

    if len(password) < 8:
        return False, "Password must contain at least 8 characters."

    if not re.search(r"[A-Z]", password):
        return False, "Password must contain at least one uppercase letter."

    if not re.search(r"[a-z]", password):
        return False, "Password must contain at least one lowercase letter."

    if not re.search(r"[0-9]", password):
        return False, "Password must contain at least one number."

    if not re.search(r"[^A-Za-z0-9]", password):
        return False, "Password must contain at least one special character."

    return True, ""

# ============================================================
# PASSWORD HASHING
# ============================================================

def hash_password(password):
    return hashlib.sha256(
        password.encode("utf-8")
    ).hexdigest()

# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(uploaded_file):

    reader = PyPDF2.PdfReader(uploaded_file)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text

# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(uploaded_file):

    document = docx.Document(uploaded_file)

    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:
                text += cell.text + "\n"

    return text

# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(uploaded_file):

    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    elif filename.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    return ""

# ============================================================
# RESUME FINGERPRINT
# ============================================================

def create_resume_fingerprint(text):

    normalized_text = text.lower()

    normalized_text = re.sub(
        r"\s+",
        " ",
        normalized_text
    )

    normalized_text = re.sub(
        r"[^a-z0-9+#.]+",
        " ",
        normalized_text
    )

    normalized_text = re.sub(
        r"\s+",
        " ",
        normalized_text
    ).strip()

    fingerprint = hashlib.sha256(
        normalized_text.encode("utf-8")
    ).hexdigest()

    return fingerprint

# ============================================================
# FIND SKILLS
# ============================================================

def find_skills(text, skills):

    found = []

    text_lower = text.lower()

    for skill in skills:

        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, text_lower):
            found.append(skill)

    return found

# ============================================================
# FIND PROGRAMMING LANGUAGES
# ============================================================

def find_programming_languages(text):

    found = []

    text_lower = text.lower()

    for language in CORE_LANGUAGES:

        if language == "C++":

            pattern = r"(?<!\w)c\+\+(?!\w)"

        else:

            pattern = (
                r"(?<!\w)"
                + re.escape(language.lower())
                + r"(?!\w)"
            )

        if re.search(pattern, text_lower):
            found.append(language)

    return found

# ============================================================
# EDUCATION
# ============================================================

EDUCATION_KEYWORDS = [
    "B.Tech",
    "B.E",
    "Bachelor",
    "M.Tech",
    "M.E",
    "Master",
    "Degree",
    "B.Sc",
    "BCA",
    "MCA"
]


def find_education(text):

    found = []

    text_lower = text.lower()

    for education in EDUCATION_KEYWORDS:

        if education.lower() in text_lower:
            found.append(education)

    return found

# ============================================================
# FIND PROJECTS
# ============================================================

def find_projects(text):

    matches = re.findall(
        r"\bprojects?\b",
        text,
        flags=re.IGNORECASE
    )

    return len(matches)

# ============================================================
# FIND EXPERIENCE
# ============================================================

def find_experience_months(text):

    total_months = 0

    year_matches = re.findall(
        r"(\d+(?:\.\d+)?)\s*(?:years?|yrs?)",
        text,
        flags=re.IGNORECASE
    )

    for value in year_matches:

        try:
            total_months += int(
                float(value) * 12
            )

        except Exception:
            pass

    month_matches = re.findall(
        r"(\d+)\s*(?:months?|mos?)",
        text,
        flags=re.IGNORECASE
    )

    for value in month_matches:

        try:
            total_months += int(value)

        except Exception:
            pass

    return total_months

# ============================================================
# JOB DESCRIPTION SKILLS
# ============================================================

def extract_required_skills_from_description(description):

    all_skills = set()

    for skills in ROLE_SKILLS.values():

        for skill in skills:
            all_skills.add(skill)

    for skill in RECOMMENDATIONS.keys():
        all_skills.add(skill)

    for language in CORE_LANGUAGES:
        all_skills.add(language)

    return find_skills(
        description,
        list(all_skills)
    )

# ============================================================
# RESUME SECTION ANALYSIS
# ============================================================

def detect_resume_sections(text):

    patterns = {
        "Contact Information": [
            r"\bemail\b",
            r"\bphone\b",
            r"\bmobile\b",
            r"\bcontact\b"
        ],

        "Professional Summary / Objective": [
            r"\bprofessional summary\b",
            r"\bcareer objective\b",
            r"\bsummary\b",
            r"\bobjective\b"
        ],

        "Education": [
            r"\beducation\b",
            r"\bacademic qualifications\b",
            r"\bacademic background\b"
        ],

        "Technical Skills": [
            r"\btechnical skills\b",
            r"\bskills\b",
            r"\bskill set\b",
            r"\btechnologies\b"
        ],

        "Projects": [
            r"\bprojects\b",
            r"\bacademic projects\b",
            r"\bpersonal projects\b"
        ],

        "Work Experience / Internships": [
            r"\bwork experience\b",
            r"\bprofessional experience\b",
            r"\binternship\b",
            r"\binternships\b",
            r"\bemployment history\b"
        ],

        "Certifications": [
            r"\bcertifications\b",
            r"\bcertificates\b",
            r"\blicenses\b"
        ],

        "Achievements": [
            r"\bachievements\b",
            r"\bawards\b",
            r"\baccomplishments\b"
        ],

        "Declaration": [
            r"\bdeclaration\b"
        ]
    }

    detected = []
    missing = []

    lower_text = text.lower()

    for section, section_patterns in patterns.items():

        if any(
            re.search(
                pattern,
                lower_text
            )
            for pattern in section_patterns
        ):
            detected.append(section)

        else:
            missing.append(section)

    return detected, missing


def calculate_resume_completeness(detected_sections):

    important = [
        "Education",
        "Technical Skills",
        "Projects",
        "Work Experience / Internships",
        "Professional Summary / Objective",
        "Certifications"
    ]

    present = [
        section
        for section in important
        if section in detected_sections
    ]

    completeness = (
        (len(present) / len(important)) * 100
        if important
        else 0
    )

    return completeness, present


def generate_resume_suggestions(
    text,
    missing_skills,
    detected_sections
):

    suggestions = []

    words = re.findall(
        r"\b[\w+#.-]+\b",
        text
    )

    count = len(words)

    if count < 150:

        suggestions.append(
            "Your resume appears short. Consider adding relevant projects, skills, education, and experience details."
        )

    elif count > 1000:

        suggestions.append(
            "Your resume is lengthy. Consider removing repeated or less relevant information and keeping it concise."
        )

    if "Professional Summary / Objective" not in detected_sections:

        suggestions.append(
            "Consider adding a brief professional summary or career objective."
        )

    if "Education" not in detected_sections:

        suggestions.append(
            "Add an Education section with your degree, institution, graduation year, and relevant details."
        )

    if "Technical Skills" not in detected_sections:

        suggestions.append(
            "Add a clearly labelled Technical Skills section with relevant languages, tools, and technologies."
        )

    if "Projects" not in detected_sections:

        suggestions.append(
            "Include relevant academic or personal projects and describe your contribution and technologies used."
        )

    if "Work Experience / Internships" not in detected_sections:

        suggestions.append(
            "If applicable, include internships, training, volunteering, or other relevant experience. Freshers can emphasize projects."
        )

    if "Certifications" not in detected_sections:

        suggestions.append(
            "If you have relevant certifications, list their names and issuing organizations."
        )

    if missing_skills:

        suggestions.append(
            "These target-role skills were not detected: "
            + ", ".join(missing_skills)
            + ". Include them only if you genuinely have them, and support them with examples."
        )

    if not re.search(
        r"\b\d+\s*%|\b\d+\s*(?:users|customers|projects|reports)\b",
        text,
        re.IGNORECASE
    ):

        suggestions.append(
            "Where accurate, add measurable outcomes such as records processed, reports created, or time saved."
        )

    verbs = [
        "developed",
        "created",
        "built",
        "implemented",
        "designed",
        "analysed",
        "analyzed",
        "managed",
        "improved",
        "automated",
        "tested",
        "integrated"
    ]

    if not any(
        re.search(
            r"\b" + re.escape(verb) + r"\b",
            text.lower()
        )
        for verb in verbs
    ):

        suggestions.append(
            "Use action verbs such as Developed, Implemented, Designed, Tested, or Analysed to describe your actual work."
        )

    return suggestions

# ============================================================
# SKILL GAP ANALYSIS
# ============================================================

def analyze_skill_gap(
    text,
    target_role=None,
    job_description=""
):

    if job_description.strip():

        required_skills = (
            extract_required_skills_from_description(
                job_description
            )
        )

    else:

        required_skills = ROLE_SKILLS.get(
            target_role,
            []
        )

    found_skills = find_skills(
        text,
        required_skills
    )

    missing_skills = [
        skill
        for skill in required_skills
        if skill not in found_skills
    ]

    if required_skills:

        match_percentage = (
            len(found_skills)
            / len(required_skills)
        ) * 100

    else:

        match_percentage = 0

    return (
        found_skills,
        missing_skills,
        required_skills,
        match_percentage
    )

# ============================================================
# CAREER RISK
# ============================================================

def calculate_career_risk(
    match_percentage,
    experience_months,
    project_count
):

    score = match_percentage

    if experience_months >= 24:

        score += 10

    elif experience_months >= 12:

        score += 5

    if project_count >= 5:

        score += 10

    elif project_count >= 3:

        score += 5

    score = min(
        score,
        100
    )

    if score >= 75:

        risk = "Low Risk"

    elif score >= 50:

        risk = "Medium Risk"

    else:

        risk = "High Risk"

    return risk, score

# ============================================================
# PARSE COMMA LIST
# ============================================================

def parse_comma_list(value):

    if not value.strip():
        return []

    return [
        item.strip()
        for item in value.split(",")
        if item.strip()
    ]

# ============================================================
# PARSE SKILLS WITH LEVELS
# ============================================================

def parse_skills_with_levels(value):

    skills = []

    if not value.strip():
        return skills

    items = value.split(",")

    for item in items:

        parts = item.split(":")

        skill = parts[0].strip()

        level = "Beginner"

        if len(parts) > 1:

            level = parts[1].strip()

        if skill:

            skills.append(
                {
                    "skill": skill,
                    "level": level
                }
            )

    return skills

# ============================================================
# GENERATE ROADMAP
# ============================================================

def generate_skill_roadmap(
    missing_skills,
    current_skills
):

    roadmap = []

    visited = set()

    current_skills_lower = {
        str(skill).lower()
        for skill in current_skills
    }

    def add_skill(skill):

        if skill in visited:
            return

        visited.add(skill)

        for dependency in SKILL_DEPENDENCIES.get(
            skill,
            []
        ):

            if dependency.lower() not in current_skills_lower:

                add_skill(dependency)

        if skill.lower() not in current_skills_lower:

            if skill not in roadmap:

                roadmap.append(skill)

    for skill in missing_skills:

        add_skill(skill)

    return roadmap

# ============================================================
# LOGIN PAGE
# ============================================================

def login_page():

    if st.session_state.get(
        "registration_success",
        False
    ):

        st.success(
            "Account created successfully! Please login to continue."
        )

        st.session_state[
            "registration_success"
        ] = False

    st.markdown(
        """
        <div style="text-align:center;">
            <h1>📊 Career Risk Analyzer</h1>
            <p>Login to continue</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    left, center, right = st.columns(
        [2.0, 2.0, 2.0]
    )

    with center:

        users = load_users()

        if not isinstance(
            users,
            dict
        ):

            users = {}

        if (
            "users" in users
            and isinstance(
                users["users"],
                dict
            )
        ):

            account_users = users["users"]

        else:

            account_users = users

        login_tab, register_tab = st.tabs(
            [
                "🔐 Login",
                "📝 Register"
            ]
        )

        # ====================================================
        # LOGIN TAB
        # ====================================================

        with login_tab:

            with st.form("login_form"):

                username = st.text_input(
                    "Username",
                    placeholder="Enter your username"
                )

                password = st.text_input(
                    "Password",
                    type="password",
                    placeholder="Enter your password"
                )

                login_button = st.form_submit_button(
                    "🔐 Login",
                    use_container_width=True
                )

            if login_button:

                username = username.strip()

                if not username:

                    st.error(
                        "Please enter your username."
                    )

                elif not password:

                    st.error(
                        "Please enter your password."
                    )

                elif username not in account_users:

                    st.error(
                        "Username does not exist."
                    )

                else:

                    user_data = account_users.get(
                        username
                    )

                    stored_password = ""

                    if isinstance(
                        user_data,
                        str
                    ):

                        stored_password = user_data

                    elif isinstance(
                        user_data,
                        dict
                    ):

                        stored_password = user_data.get(
                            "password",
                            ""
                        )

                    password_is_correct = False

                    if stored_password == hash_password(
                        password
                    ):

                        password_is_correct = True

                    elif stored_password == password:

                        password_is_correct = True

                    if password_is_correct:

                        if stored_password == password:

                            account_users[
                                username
                            ] = hash_password(
                                password
                            )

                            if (
                                "users" in users
                                and isinstance(
                                    users["users"],
                                    dict
                                )
                            ):

                                users["users"] = account_users

                            else:

                                users = account_users

                            save_users(users)

                        st.session_state[
                            "username"
                        ] = username

                        st.session_state[
                            "page"
                        ] = "analyzer"

                        st.session_state[
                            "selected_history_index"
                        ] = None

                        st.session_state[
                            "open_roadmap_skill"
                        ] = None

                        st.session_state[
                            "clear_history_confirm"
                        ] = False

                        st.session_state[
                            "registration_success"
                        ] = False

                        st.success(
                            "Login successful!"
                        )

                        st.rerun()

                    else:

                        st.error(
                            "Incorrect password."
                        )

        # ====================================================
        # REGISTER TAB
        # ====================================================

        with register_tab:

            with st.form("register_form"):

                new_username = st.text_input(
                    "Username",
                    placeholder="Create a username"
                )

                new_password = st.text_input(
                    "Password",
                    type="password",
                    placeholder="Create a password"
                )

                st.markdown(
                    """
                    <p style="
                        color: #9aa0a6;
                        font-size: 12px;
                        margin-top: -8px;
                        margin-bottom: 12px;
                    ">
                        Minimum 8 characters • 1 uppercase • 1 lowercase • 1 number • 1 special character
                    </p>
                    """,
                    unsafe_allow_html=True
                )

                confirm_password = st.text_input(
                    "Confirm Password",
                    type="password",
                    placeholder="Re-enter your password"
                )

                register_button = st.form_submit_button(
                    "📝 Create Account",
                    use_container_width=True
                )

            if register_button:

                username = new_username.strip()

                if not username:

                    st.error(
                        "Please enter a username."
                    )

                elif username in account_users:

                    st.error(
                        "Username already exists."
                    )

                elif not new_password:

                    st.error(
                        "Please enter a password."
                    )

                else:

                    valid, message = validate_password(
                        new_password
                    )

                    if not valid:

                        st.error(
                            message
                        )

                    elif new_password != confirm_password:

                        st.error(
                            "Passwords do not match."
                        )

                    else:

                        hashed_password = hash_password(
                            new_password
                        )

                        if (
                            "users" in users
                            and isinstance(
                                users["users"],
                                dict
                            )
                        ):

                            users["users"][username] = (
                                hashed_password
                            )

                        else:

                            users[username] = (
                                hashed_password
                            )

                        save_users(users)

                        st.session_state[
                            "registration_success"
                        ] = True

                        st.session_state[
                            "page"
                        ] = "login"

                        st.session_state[
                            "username"
                        ] = ""

                        st.session_state[
                            "selected_history_index"
                        ] = None

                        st.session_state[
                            "open_roadmap_skill"
                        ] = None

                        st.session_state[
                            "clear_history_confirm"
                        ] = False

                        st.rerun()

# ============================================================
# DASHBOARD PAGE
# ============================================================

def dashboard_page():

    username = st.session_state["username"]

    profiles = load_profiles()
    history_data = load_history()

    profile = profiles.get(
        username,
        {}
    )

    if not isinstance(
        profile,
        dict
    ):

        profile = {}

    if not isinstance(
        history_data,
        dict
    ):

        history_data = {}

    user_history = history_data.get(
        username,
        []
    )

    if not isinstance(
        user_history,
        list
    ):

        user_history = []

    st.title("📊 Dashboard")

    st.write(
        "Welcome, " + username
    )

    st.markdown("---")

    profile_skills_data = profile.get(
        "skills",
        []
    )

    profile_skill_names = []

    for item in profile_skills_data:

        if isinstance(
            item,
            dict
        ):

            skill_name = item.get(
                "skill",
                ""
            )

            if skill_name:

                profile_skill_names.append(
                    skill_name
                )

        elif isinstance(
            item,
            str
        ):

            profile_skill_names.append(
                item
            )

    certifications = profile.get(
        "certifications",
        []
    )

    if not isinstance(
        certifications,
        list
    ):

        certifications = []

    missing_skills = []

    if user_history:

        latest = user_history[0]

        missing_skills = latest.get(
            "missing_skills",
            []
        )

    col1, col2, col3, col4 = st.columns(4)

    with col1:

        st.metric(
            "Skills Available",
            len(profile_skill_names)
        )

    with col2:

        st.metric(
            "Missing Skills",
            len(missing_skills)
        )

    with col3:

        st.metric(
            "Certifications",
            len(certifications)
        )

    with col4:

        st.metric(
            "Assessments",
            len(user_history)
        )

    st.markdown("---")

    st.subheader(
        "👤 Career Profile"
    )

    profile_col1, profile_col2 = st.columns(2)

    with profile_col1:

        current_role = profile.get(
            "current_role",
            ""
        )

        target_role = profile.get(
            "target_job_role",
            ""
        )

        experience = profile.get(
            "experience",
            ""
        )

        projects = profile.get(
            "projects",
            ""
        )

        st.write(
            "**Current Role:** "
            + (
                current_role
                if current_role
                else "Not added"
            )
        )

        st.write(
            "**Target Job Role:** "
            + (
                target_role
                if target_role
                else "Not added"
            )
        )

        st.write(
            "**Experience:** "
            + (
                str(experience)
                if experience
                else "Not added"
            )
        )

        st.write(
            "**Projects:** "
            + (
                str(projects)
                if projects
                else "Not added"
            )
        )

    with profile_col2:

        education = profile.get(
            "education",
            ""
        )

        st.write(
            "**Education:** "
            + (
                education
                if education
                else "Not added"
            )
        )

        st.write(
            "**Skills Available:**"
        )

        if profile_skill_names:

            st.write(
                ", ".join(
                    profile_skill_names
                )
            )

        else:

            st.write(
                "No skills added."
            )

        st.write(
            "**Certifications:**"
        )

        if certifications:

            st.write(
                ", ".join(
                    certifications
                )
            )

        else:

            st.write(
                "No certifications added."
            )

    st.markdown("---")

    st.subheader(
        "⚠️ Missing Skills"
    )

    if missing_skills:

        st.write(
            "Skills identified as missing in your latest assessment:"
        )

        st.write(
            ", ".join(
                missing_skills
            )
        )

    else:

        st.info(
            "No missing skills are available yet. Complete a resume assessment."
        )

    st.markdown("---")

    st.subheader(
        "🏆 Certifications"
    )

    if certifications:

        for certification in certifications:

            st.write(
                "• " + str(certification)
            )

    else:

        st.info(
            "No certifications added to your career profile."
        )

    st.markdown("---")

    st.subheader(
        "📈 Latest Assessment"
    )

    if user_history:

        latest = user_history[0]

        c1, c2, c3 = st.columns(3)

        with c1:

            st.metric(
                "Risk",
                latest.get(
                    "risk",
                    "N/A"
                )
            )

        with c2:

            st.metric(
                "Score",
                str(
                    round(
                        latest.get(
                            "score",
                            0
                        ),
                        1
                    )
                )
            )

        with c3:

            st.metric(
                "Skill Match",
                str(
                    round(
                        latest.get(
                            "match_percentage",
                            0
                        ),
                        1
                    )
                )
                + "%"
            )

    else:

        st.info(
            "No assessment completed yet."
        )

    st.markdown("---")

    st.subheader(
        "🔎 What the Career Risk Analyzer Does"
    )

    col1, col2, col3, col4 = st.columns(4)

    with col1:

        st.info(
            "**Resume Analysis**\n\n"
            "Extracts skills, education, programming languages, projects and experience."
        )

    with col2:

        st.info(
            "**Skill Gap Analysis**\n\n"
            "Compares your current skills with skills required for the target role."
        )

    with col3:

        st.info(
            "**Career Risk Analysis**\n\n"
            "Generates a simple career risk score based on skills, experience and projects."
        )

    with col4:

        st.info(
            "**Personalized Roadmap**\n\n"
            "Provides learning steps for missing skills."
        )

# ============================================================
# CAREER PROFILE PAGE
# ============================================================

def career_profile_page():

    username = st.session_state["username"]

    profiles = load_profiles()

    profile = profiles.get(
        username,
        {}
    )

    if not isinstance(
        profile,
        dict
    ):

        profile = {}

    st.title(
        "👤 Career Profile"
    )

    st.write(
        "Enter your current career information."
    )

    name = st.text_input(
        "Full Name",
        value=profile.get(
            "name",
            ""
        ),
        placeholder="Enter your full name, e.g. Lohitha Paduri"
    )

    current_role = st.text_input(
        "Current Role",
        value=profile.get(
            "current_role",
            ""
        ),
        placeholder="Enter your current role, e.g. Student, Intern, Software Developer"
    )

    education = st.text_input(
        "Education",
        value=profile.get(
            "education",
            ""
        ),
        placeholder="Enter your education, e.g. B.Tech Computer Science"
    )

    experience = st.text_input(
        "Experience",
        value=profile.get(
            "experience",
            ""
        ),
        placeholder="Enter your experience, e.g. Fresher, 1 year, 2 years"
    )

    skills_text = st.text_input(
        "Skills with levels",
        value=profile.get(
            "skills_text",
            ""
        ),
        placeholder="e.g. Python: Advanced, SQL: Intermediate, Excel: Beginner"
    )

    certifications_text = st.text_input(
        "Certifications",
        value=", ".join(
            profile.get(
                "certifications",
                []
            )
        ),
        placeholder="e.g. Python Certificate, Web Development Certificate"
    )

    projects = st.text_input(
        "Projects",
        value=profile.get(
            "projects",
            ""
        ),
        placeholder="Enter your projects, e.g. Career Risk Analyzer, Alzheimer Detection"
    )

    interests_text = st.text_input(
        "Career Interests",
        value=", ".join(
            profile.get(
                "interests",
                []
            )
        ),
        placeholder="e.g. Data Analyst, Web Developer, Data Scientist"
    )

    target_roles = [
        "Select target role",
        "Data Analyst",
        "Software Developer",
        "Data Scientist",
        "Machine Learning Engineer",
        "Web Developer"
    ]

    saved_target_role = profile.get(
        "target_job_role",
        "Select target role"
    )

    if saved_target_role not in target_roles:

        saved_target_role = "Select target role"

    target_job_role = st.selectbox(
        "Target Job Role",
        target_roles,
        index=target_roles.index(
            saved_target_role
        )
    )

    save_button = st.button(
        "💾 Save Career Profile",
        use_container_width=True
    )

    if save_button:

        skills = parse_skills_with_levels(
            skills_text
        )

        certifications = parse_comma_list(
            certifications_text
        )

        interests = parse_comma_list(
            interests_text
        )

        profiles[username] = {
            "name": name,
            "current_role": current_role,
            "education": education,
            "experience": experience,
            "skills_text": skills_text,
            "skills": skills,
            "certifications": certifications,
            "projects": projects,
            "interests": interests,
            "target_job_role": (
                ""
                if target_job_role == "Select target role"
                else target_job_role
            )
        }

        save_profiles(profiles)

        st.success(
            "Career profile saved successfully!"
        )

# ============================================================
# RESUME ASSESSMENT PAGE
# ============================================================

def resume_assessment_page():

    username = st.session_state["username"]

    profiles = load_profiles()

    profile = profiles.get(
        username,
        {}
    )

    if not isinstance(
        profile,
        dict
    ):

        profile = {}

    st.title(
        "📄 Resume Assessment"
    )

    st.write(
        "Upload your resume and compare your skills with the required skills."
    )

    uploaded_file = st.file_uploader(
        "Upload Resume",
        type=["pdf", "docx"]
    )

    assessment_type = st.radio(
        "Assessment Type",
        [
            "Job Role",
            "Job Description"
        ],
        horizontal=True
    )

    target_role = None
    job_description = ""

    if assessment_type == "Job Role":

        st.write(
            "Select the job role you want to assess your resume against."
        )

        target_role = st.selectbox(
            "Select Target Job Role",
            list(ROLE_SKILLS.keys())
        )

    else:

        job_description = st.text_area(
            "Paste Job Description",
            placeholder="Paste the complete job description here..."
        )

    analyze_button = st.button(
        "🔍 Analyze Resume",
        use_container_width=True
    )

    if analyze_button:

        if uploaded_file is None:

            st.warning(
                "Please upload a PDF or DOCX resume."
            )

            return

        try:

            resume_text = extract_resume_text(
                uploaded_file
            )

        except Exception as error:

            st.error(
                "Unable to read the uploaded resume."
            )

            st.exception(
                error
            )

            return

        if not resume_text.strip():

            st.error(
                "No readable text was found in the resume. Scanned PDFs may require OCR."
            )

            return

        fingerprint = create_resume_fingerprint(
            resume_text
        )

        (
            found_skills,
            missing_skills,
            required_skills,
            match_percentage
        ) = analyze_skill_gap(
            resume_text,
            target_role,
            job_description
        )

        education = find_education(
            resume_text
        )

        programming_languages = find_programming_languages(
            resume_text
        )

        project_count = find_projects(
            resume_text
        )

        experience_months = find_experience_months(
            resume_text
        )

        risk, score = calculate_career_risk(
            match_percentage,
            experience_months,
            project_count
        )

        st.markdown("---")

        st.subheader(
            "📊 Assessment Result"
        )

        result_col1, result_col2, result_col3 = st.columns(3)

        with result_col1:

            st.metric(
                "Career Risk",
                risk
            )

        with result_col2:

            st.metric(
                "Risk Score",
                str(
                    round(
                        score,
                        1
                    )
                )
            )

        with result_col3:

            st.metric(
                "Skill Match",
                str(
                    round(
                        match_percentage,
                        1
                    )
                )
                + "%"
            )

        st.markdown("---")

        info1, info2, info3, info4 = st.columns(4)

        with info1:

            st.metric(
                "Skills Detected",
                len(found_skills)
            )

        with info2:

            st.metric(
                "Education",
                len(education)
            )

        with info3:

            st.metric(
                "Projects",
                project_count
            )

        with info4:

            st.metric(
                "Experience",
                str(
                    experience_months
                )
                + " months"
            )

        st.markdown("---")

        st.subheader(
            "✅ Skills Available"
        )

        if found_skills:

            st.write(
                ", ".join(
                    found_skills
                )
            )

        else:

            st.info(
                "No required skills were detected."
            )

        st.subheader(
            "⚠️ Missing Skills"
        )

        if missing_skills:

            st.write(
                ", ".join(
                    missing_skills
                )
            )

        else:

            st.success(
                "No missing skills detected."
            )

        st.subheader(
            "💻 Programming Languages"
        )

        if programming_languages:

            st.write(
                ", ".join(
                    programming_languages
                )
            )

        else:

            st.info(
                "No programming languages detected."
            )

        # ====================================================
        # SECTION DETECTION AND COMPLETENESS
        # ====================================================

        st.markdown("---")

        st.subheader(
            "🧩 Resume Section Analysis"
        )

        detected_sections, missing_sections = detect_resume_sections(
            resume_text
        )

        completeness, detected_important_sections = calculate_resume_completeness(
            detected_sections
        )

        st.write(
            "**Resume Section Completeness**"
        )

        st.progress(
            int(completeness)
        )

        st.write(
            str(
                round(
                    completeness,
                    1
                )
            )
            + "% of the checked important sections were detected."
        )

        section_col1, section_col2 = st.columns(2)

        with section_col1:

            st.markdown(
                "#### ✅ Detected Sections"
            )

            if detected_sections:

                for section in detected_sections:

                    st.write(
                        "• " + section
                    )

            else:

                st.info(
                    "No standard resume section headings were detected."
                )

        with section_col2:

            st.markdown(
                "#### ⚠️ Sections Not Detected"
            )

            if missing_sections:

                for section in missing_sections:

                    st.write(
                        "• " + section
                    )

            else:

                st.success(
                    "All checked section headings were detected."
                )

        st.caption(
            "Detection uses keywords and headings. A section may exist but not be detected if it uses different wording."
        )

        # ====================================================
        # DETAILED SKILL MATCH
        # ====================================================

        st.markdown("---")

        st.subheader(
            "🎯 Detailed Skill Match"
        )

        if required_skills:

            skill_col1, skill_col2, skill_col3 = st.columns(3)

            with skill_col1:

                st.metric(
                    "Required Skills",
                    len(required_skills)
                )

            with skill_col2:

                st.metric(
                    "Matched Skills",
                    len(found_skills)
                )

            with skill_col3:

                st.metric(
                    "Missing Skills",
                    len(missing_skills)
                )

            st.write(
                "**Skill Match Progress**"
            )

            st.progress(
                int(match_percentage)
            )

            st.write(
                str(
                    round(
                        match_percentage,
                        1
                    )
                )
                + "% of identified required skills were detected."
            )

            with st.expander(
                "View Matched Skills"
            ):

                if found_skills:

                    for skill in found_skills:

                        st.write(
                            "✅ " + skill
                        )

                else:

                    st.info(
                        "No matching skills were detected."
                    )

            with st.expander(
                "View Missing Skills"
            ):

                if missing_skills:

                    for skill in missing_skills:

                        st.write(
                            "⚠️ " + skill
                        )

                else:

                    st.success(
                        "No missing skills were detected from the selected requirements."
                    )

        else:

            st.info(
                "No required skills were identified, so a meaningful skill match percentage cannot be calculated."
            )

        # ====================================================
        # PERSONALIZED IMPROVEMENT SUGGESTIONS
        # ====================================================

        st.markdown("---")

        st.subheader(
            "💡 Personalized Resume Improvement Suggestions"
        )

        resume_suggestions = generate_resume_suggestions(
            resume_text,
            missing_skills,
            detected_sections
        )

        if resume_suggestions:

            for suggestion_index, suggestion in enumerate(
                resume_suggestions,
                start=1
            ):

                st.write(
                    str(suggestion_index)
                    + ". "
                    + suggestion
                )

        else:

            st.success(
                "No additional suggestions were generated by these checks."
            )

        # ====================================================
        # SAVE ASSESSMENT HISTORY
        # ====================================================

        history_data = load_history()

        if not isinstance(
            history_data,
            dict
        ):

            history_data = {}

        if username not in history_data:

            history_data[username] = []

        assessment = {

            "date": datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),

            "assessment_type": assessment_type,

            "target_role": (
                target_role
                if target_role
                else "Job Description"
            ),

            "risk": risk,

            "score": score,

            "match_percentage": match_percentage,

            "required_skills": required_skills,

            "found_skills": found_skills,

            "missing_skills": missing_skills,

            "education": education,

            "programming_languages": programming_languages,

            "project_count": project_count,

            "experience_months": experience_months,

            "certifications": profile.get(
                "certifications",
                []
            ),

            "resume_name": uploaded_file.name,

            "resume_fingerprint": fingerprint
        }

        history_data[username].insert(
            0,
            assessment
        )

        save_history(
            history_data
        )

        st.success(
            "Assessment saved to your history."
        )

# ============================================================
# ROADMAP PAGE
# ============================================================

def roadmap_page():

    username = st.session_state["username"]

    history_data = load_history()

    if not isinstance(
        history_data,
        dict
    ):

        history_data = {}

    user_history = history_data.get(
        username,
        []
    )

    if not isinstance(
        user_history,
        list
    ):

        user_history = []

    st.title(
        "🗺️ Personalized Skill Roadmap"
    )

    st.write(
        "Learn the missing skills step by step using concept-level learning resources."
    )

    if not user_history:

        st.info(
            "Complete a resume assessment to generate your roadmap."
        )

        return

    unique_resumes = []

    seen_fingerprints = set()

    for assessment in user_history:

        fingerprint = assessment.get(
            "resume_fingerprint",
            ""
        )

        if not fingerprint:

            normalized_name = re.sub(
                r"\s+",
                " ",
                assessment.get(
                    "resume_name",
                    ""
                ).lower().strip()
            )

            fingerprint = (
                "filename:"
                + normalized_name
            )

        if fingerprint not in seen_fingerprints:

            seen_fingerprints.add(
                fingerprint
            )

            unique_resumes.append(
                assessment
            )

    for resume_index, assessment in enumerate(
        unique_resumes
    ):

        resume_name = assessment.get(
            "resume_name",
            "Resume"
        )

        role = assessment.get(
            "target_role",
            "Job Description"
        )

        missing_skills = assessment.get(
            "missing_skills",
            []
        )

        found_skills = assessment.get(
            "found_skills",
            []
        )

        roadmap = generate_skill_roadmap(
            missing_skills,
            found_skills
        )

        st.subheader(
            "📄 " + resume_name
        )

        st.write(
            "**Target Role:** "
            + str(role)
        )

        if not roadmap:

            st.success(
                "No additional roadmap skills are required."
            )

        else:

            st.write(
                "The roadmap below shows the skills you need to learn."
            )

            for skill_index, skill in enumerate(
                roadmap
            ):

                left_space, skill_center, right_space = st.columns(
                    [1, 5, 1]
                )

                with skill_center:

                    if st.button(
                        "📚 " + skill,
                        key=(
                            "roadmap_skill_"
                            + str(resume_index)
                            + "_"
                            + str(skill_index)
                            + "_"
                            + skill
                        ),
                        use_container_width=True
                    ):

                        if (
                            st.session_state[
                                "open_roadmap_skill"
                            ]
                            == skill
                        ):

                            st.session_state[
                                "open_roadmap_skill"
                            ] = None

                        else:

                            st.session_state[
                                "open_roadmap_skill"
                            ] = skill

                        st.rerun()

                if (
                    st.session_state[
                        "open_roadmap_skill"
                    ]
                    == skill
                ):

                    details = ROADMAP_DETAILS.get(
                        skill,
                        {}
                    )

                    detail_left, detail_center, detail_right = st.columns(
                        [1, 4, 1]
                    )

                    with detail_center:

                        st.markdown(
                            "### 📚 " + skill
                        )

                        st.markdown(
                            "#### 1️⃣ Concepts You Need to Learn"
                        )

                        concepts = details.get(
                            "concepts",
                            []
                        )

                        if concepts:

                            for concept_index, concept in enumerate(
                                concepts,
                                start=1
                            ):

                                st.write(
                                    str(concept_index)
                                    + ". "
                                    + concept["name"]
                                )

                                st.markdown(
                                    "🔗 [Learn this concept]("
                                    + concept["link"]
                                    + ")"
                                )

                        else:

                            st.write(
                                RECOMMENDATIONS.get(
                                    skill,
                                    "Learn the fundamentals of this skill."
                                )
                            )

                        st.markdown(
                            "#### 2️⃣ Practice"
                        )

                        st.write(
                            details.get(
                                "practice",
                                "Practice this skill using real-world examples and exercises."
                            )
                        )

                        st.markdown(
                            "#### 3️⃣ Real-Time Project"
                        )

                        st.write(
                            details.get(
                                "project",
                                "Build a real-world project using this skill."
                            )
                        )

                        st.markdown(
                            "#### 4️⃣ Main Learning Resource"
                        )

                        resource = LEARNING_RESOURCES.get(
                            skill
                        )

                        if resource:

                            st.markdown(
                                "[Open complete learning resource]("
                                + resource
                                + ")"
                            )

                        st.markdown("---")

# ============================================================
# HISTORY DETAILS
# ============================================================

def display_history_details(
    assessment
):

    st.markdown(
        "### 📋 Assessment Details"
    )

    c1, c2, c3 = st.columns(3)

    with c1:

        st.write(
            "**Risk:**",
            assessment.get(
                "risk",
                "N/A"
            )
        )

    with c2:

        st.write(
            "**Score:**",
            round(
                assessment.get(
                    "score",
                    0
                ),
                1
            )
        )

    with c3:

        st.write(
            "**Skill Match:**",
            str(
                round(
                    assessment.get(
                        "match_percentage",
                        0
                    ),
                    1
                )
            )
            + "%"
        )

    st.write(
        "**Assessment Type:**",
        assessment.get(
            "assessment_type",
            "N/A"
        )
    )

    st.write(
        "**Target Role:**",
        assessment.get(
            "target_role",
            "N/A"
        )
    )

    st.write(
        "**Resume:**",
        assessment.get(
            "resume_name",
            "N/A"
        )
    )

    st.write(
        "**Date:**",
        assessment.get(
            "date",
            "N/A"
        )
    )

    st.write(
        "**Skills Available:**"
    )

    found_skills = assessment.get(
        "found_skills",
        []
    )

    if found_skills:

        st.write(
            ", ".join(
                found_skills
            )
        )

    else:

        st.write(
            "None detected"
        )

    st.write(
        "**Missing Skills:**"
    )

    missing_skills = assessment.get(
        "missing_skills",
        []
    )

    if missing_skills:

        st.write(
            ", ".join(
                missing_skills
            )
        )

    else:

        st.write(
            "None"
        )

    st.write(
        "**Education:**"
    )

    education = assessment.get(
        "education",
        []
    )

    if education:

        st.write(
            ", ".join(
                education
            )
        )

    else:

        st.write(
            "None detected"
        )

    st.write(
        "**Programming Languages:**"
    )

    programming_languages = assessment.get(
        "programming_languages",
        []
    )

    if programming_languages:

        st.write(
            ", ".join(
                programming_languages
            )
        )

    else:

        st.write(
            "None detected"
        )

    st.write(
        "**Projects:**",
        assessment.get(
            "project_count",
            0
        )
    )

    st.write(
        "**Experience:**",
        str(
            assessment.get(
                "experience_months",
                0
            )
        )
        + " months"
    )

    certifications = assessment.get(
        "certifications",
        []
    )

    st.write(
        "**Certifications:**"
    )

    if certifications:

        st.write(
            ", ".join(
                certifications
            )
        )

    else:

        st.write(
            "None"
        )

# ============================================================
# HISTORY PAGE
# ============================================================

def history_page():

    username = st.session_state["username"]

    history_data = load_history()

    if not isinstance(
        history_data,
        dict
    ):

        history_data = {}

    user_history = history_data.get(
        username,
        []
    )

    if not isinstance(
        user_history,
        list
    ):

        user_history = []

    st.title(
        "📜 Assessment History"
    )

    if not user_history:

        st.info(
            "No previous assessments found."
        )

    else:

        for index, assessment in enumerate(
            user_history
        ):

            resume_name = assessment.get(
                "resume_name",
                "Resume"
            )

            role = assessment.get(
                "target_role",
                "Job Description"
            )

            date = assessment.get(
                "date",
                ""
            )

            risk = assessment.get(
                "risk",
                ""
            )

            match = assessment.get(
                "match_percentage",
                0
            )

            label = (
                "📄 "
                + resume_name
                + " | "
                + str(role)
                + " | "
                + str(date)
                + " | "
                + str(risk)
                + " | Match: "
                + str(
                    round(
                        match,
                        1
                    )
                )
                + "%"
            )

            if st.button(
                label,
                key="history_" + str(index),
                use_container_width=True
            ):

                if (
                    st.session_state[
                        "selected_history_index"
                    ]
                    == index
                ):

                    st.session_state[
                        "selected_history_index"
                    ] = None

                else:

                    st.session_state[
                        "selected_history_index"
                    ] = index

                st.rerun()

            if (
                st.session_state[
                    "selected_history_index"
                ]
                == index
            ):

                display_history_details(
                    assessment
                )

    st.markdown("---")

    if user_history:

        if not st.session_state[
            "clear_history_confirm"
        ]:

            if st.button(
                "🗑️ Clear History",
                use_container_width=True
            ):

                st.session_state[
                    "clear_history_confirm"
                ] = True

        else:

            st.warning(
                "Are you sure you want to clear all assessment history?"
            )

            confirm_col1, confirm_col2 = st.columns(2)

            with confirm_col1:

                if st.button(
                    "Yes, Clear History",
                    use_container_width=True
                ):

                    history_data[
                        username
                    ] = []

                    save_history(
                        history_data
                    )

                    st.session_state[
                        "selected_history_index"
                    ] = None

                    st.session_state[
                        "clear_history_confirm"
                    ] = False

                    st.success(
                        "Assessment history cleared."
                    )

                    st.rerun()

            with confirm_col2:

                if st.button(
                    "Cancel",
                    use_container_width=True
                ):

                    st.session_state[
                        "clear_history_confirm"
                    ] = False

                    st.rerun()

# ============================================================
# MAIN ANALYZER PAGE
# ============================================================

def analyzer_page():

    username = st.session_state[
        "username"
    ]

    st.sidebar.title(
        "📊 Career Risk Analyzer"
    )

    st.sidebar.write(
        "👋 Welcome, " + username + "!"
    )

    st.sidebar.success(
        "You are successfully logged in."
    )

    selected_page = st.sidebar.radio(
        "Navigation",
        [
            "Dashboard",
            "Career Profile",
            "Resume Assessment",
            "🗺️ Roadmap",
            "📜 History"
        ],
        key="navigation_page"
    )

    st.sidebar.markdown("---")

    if st.sidebar.button(
        "Logout",
        use_container_width=True
    ):

        st.session_state[
            "username"
        ] = ""

        st.session_state[
            "page"
        ] = "login"

        st.session_state[
            "selected_history_index"
        ] = None

        st.session_state[
            "open_roadmap_skill"
        ] = None

        st.session_state[
            "clear_history_confirm"
        ] = False

        st.session_state[
            "registration_success"
        ] = False

        st.rerun()

    if selected_page == "Dashboard":

        dashboard_page()

    elif selected_page == "Career Profile":

        career_profile_page()

    elif selected_page == "Resume Assessment":

        resume_assessment_page()

    elif selected_page == "🗺️ Roadmap":

        roadmap_page()

    elif selected_page == "📜 History":

        history_page()

# ============================================================
# START APPLICATION
# ============================================================

if st.session_state["page"] == "login":

    login_page()

else:

    analyzer_page()